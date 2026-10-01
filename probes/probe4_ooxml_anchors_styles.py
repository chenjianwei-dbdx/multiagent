"""Probe 4 — python-docx/lxml 对 OOXML 锚点与样式的读取能力(离线)。

目的:验证 P1 所需的五类 OOXML 读取路径在 python-docx 1.2.0 + lxml 6.1.3 下是否可行、
哪些必须绕过 python-docx 高层 API 直接用 lxml 解析 zip 内 XML:

1. 枚举 w:bookmarkStart 的 name/id 及所在段落;
2. 枚举 w:t 文本及其 run/段落结构,识别 w:tab/w:br(含表格单元格内段落);
3. 解析 run 有效样式:direct formatting(w:rPr) → 段落样式 → styles.xml basedOn 链 →
   docDefaults;取出字体(w:rFonts ascii/eastAsia,注意 *Theme 间接)与字号(w:sz 半磅);
4. numbering.xml:numId → abstractNumId → 级别定义(numFmt/lvlText);
5. 页眉页脚 part 枚举(word/header*.xml)及 sectPr 的 headerReference 关联。

文档侧构造:python-docx 建 bookmark(底层 lxml 注入)、普通段落、标题段落、List Bullet
列表段、2x2 表格(单元格放占位符段落)、带 direct formatting 的 run(含 eastAsia 字体)、
页眉页脚文本。"高层 API 不便、需直接 lxml"的结论记入 findings。
"""

from __future__ import annotations

import io
import itertools
import posixpath
import zipfile
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

BOOKMARK_NAME = "omas_probe_anchor"
BOOKMARK_ID = "137"
DIRECT_FONT_ASCII = "Arial"
DIRECT_FONT_EAST_ASIA = "SimSun"
DIRECT_SIZE_PT = 14.0


def _build_docx_bytes() -> bytes:
    """构造包含全部受测结构的 DOCX 字节。"""
    doc = Document()

    # 1) 带受控锚点的普通段落(bookmarkStart/End 用底层 lxml 注入,包住 run)
    p_anchor = doc.add_paragraph("anchor paragraph text")
    r_anchor = p_anchor.runs[0]._r
    bs = p_anchor._p.makeelement(
        qn("w:bookmarkStart"), {qn("w:id"): BOOKMARK_ID, qn("w:name"): BOOKMARK_NAME}
    )
    be = p_anchor._p.makeelement(qn("w:bookmarkEnd"), {qn("w:id"): BOOKMARK_ID})
    r_anchor.addprevious(bs)
    r_anchor.addnext(be)

    # 2) 标题样式段落
    doc.add_paragraph("heading text", style="Heading 1")

    # 3) 列表段落(List Bullet 样式引用 numbering)
    doc.add_paragraph("bullet item", style="List Bullet")

    # 4) direct formatting run:西文字体 + 中文字体 + 字号
    p_direct = doc.add_paragraph()
    run = p_direct.add_run("direct formatted 中文 run")
    run.font.name = DIRECT_FONT_ASCII
    run.font.size = Pt(DIRECT_SIZE_PT)
    rpr = run._r.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), DIRECT_FONT_EAST_ASIA)

    # 5) 2x2 表格,右下单元格放占位符段落
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(1, 1).paragraphs[0].text = "{{ payload }}"

    # 6) 页眉页脚
    doc.sections[0].header.paragraphs[0].text = "HEADER TEXT"
    doc.sections[0].footer.paragraphs[0].text = "FOOTER TEXT"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _para_text(p_el: Any) -> str:
    """段落逻辑文本:w:t 原文 + w:tab→\\t + w:br→\\n(按文档顺序)。"""
    parts: list[str] = []
    for el in p_el.iter():
        tag = etree.QName(el).localname
        if tag == "t":
            parts.append(el.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag == "br":
            parts.append("\n")
    return "".join(parts)


def _w(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _list_paragraphs(root: Any) -> list[Any]:
    body = root.find(_w("body"))
    paras: list[Any] = []
    for p in body.iter(_w("p")):
        paras.append(p)
    return paras


def _enum_bookmarks(doc_root: Any) -> list[dict[str, Any]]:
    """路径 1:所有 w:bookmarkStart 的 name/id + 所在段落索引与文本。"""
    paras = _list_paragraphs(doc_root)
    out = []
    for bs in doc_root.iter(_w("bookmarkStart")):
        # 所在段落:沿祖先找 w:p
        ancestor_p = bs.getparent()
        while ancestor_p is not None and etree.QName(ancestor_p).localname != "p":
            ancestor_p = ancestor_p.getparent()
        idx = paras.index(ancestor_p) if ancestor_p in paras else None
        has_matching_end = any(
            be.get(_w("id")) == bs.get(_w("id")) for be in doc_root.iter(_w("bookmarkEnd"))
        )
        out.append(
            {
                "name": bs.get(_w("name")),
                "id": bs.get(_w("id")),
                "paragraph_index": idx,
                "paragraph_text": _para_text(ancestor_p) if ancestor_p is not None else None,
                "has_matching_bookmarkEnd": has_matching_end,
            }
        )
    return out


def _enum_text_structure(doc_root: Any) -> dict[str, Any]:
    """路径 2:w:t/w:tab/w:br 的 run/段落结构枚举(含表格内段落与单元格坐标)。"""
    paras = _list_paragraphs(doc_root)
    body = doc_root.find(_w("body"))

    para_records = []
    for i, p in enumerate(paras):
        runs = []
        for r_el in p.findall(_w("r")):
            run_text = ""
            n_t = n_tab = n_br = 0
            for child in r_el:
                tag = etree.QName(child).localname
                if tag == "t":
                    run_text += child.text or ""
                    n_t += 1
                elif tag == "tab":
                    run_text += "\t"
                    n_tab += 1
                elif tag == "br":
                    run_text += "\n"
                    n_br += 1
            runs.append({"text": run_text, "n_wt": n_t, "n_tab": n_tab, "n_br": n_br})
        para_records.append({"index": i, "runs": runs})

    # 表格:行/列 → 单元格内段落文本
    tables = []
    for tbl_el in body.iter(_w("tbl")):
        rows = []
        for ri, tr in enumerate(tbl_el.findall(_w("tr"))):
            cells = []
            for _ci, tc in enumerate(tr.findall(_w("tc"))):
                cells.append([_para_text(p) for p in tc.findall(_w("p"))])
            rows.append({"row": ri, "cells": cells})
        tables.append({"rows": rows})

    return {
        "paragraphs": para_records,
        "tables": tables,
        "total_wt": len(list(doc_root.iter(_w("t")))),
        "total_tab": len(list(doc_root.iter(_w("tab")))),
        "total_br": len(list(doc_root.iter(_w("br")))),
    }


def _rfonts_attrs(rfonts: Any) -> dict[str, Any]:
    """提取 w:rFonts 的字体属性(ascii/eastAsia 与 *Theme 间接引用)。"""
    if rfonts is None:
        return {
            "rFonts_ascii": None,
            "rFonts_eastAsia": None,
            "rFonts_asciiTheme": None,
            "rFonts_eastAsiaTheme": None,
        }
    return {
        "rFonts_ascii": rfonts.get(_w("ascii")),
        "rFonts_eastAsia": rfonts.get(_w("eastAsia")),
        "rFonts_asciiTheme": rfonts.get(_w("asciiTheme")),
        "rFonts_eastAsiaTheme": rfonts.get(_w("eastAsiaTheme")),
    }


def _style_chain(styles_root: Any, style_id: str) -> list[dict[str, Any]]:
    """styles.xml 中 style 的 basedOn 链(含每级的 rPr 摘要)。"""
    chain = []
    current = style_id
    seen: set[str] = set()
    while current and current not in seen:
        seen.add(current)
        el = styles_root.find(f"{_w('style')}[@{_w('styleId')}='{current}']")
        if el is None:
            chain.append({"styleId": current, "found": False})
            break
        based_on = el.find(_w("basedOn"))
        sz = el.find(f"{_w('rPr')}/{_w('sz')}")
        entry: dict[str, Any] = {
            "styleId": current,
            "found": True,
            "type": el.get(_w("type")),
            "basedOn": based_on.get(_w("val")) if based_on is not None else None,
            "sz_half_points": sz.get(_w("val")) if sz is not None else None,
        }
        entry.update(_rfonts_attrs(el.find(f"{_w('rPr')}/{_w('rFonts')}")))
        chain.append(entry)
        current = based_on.get(_w("val")) if based_on is not None else None
    return chain


def _doc_defaults(styles_root: Any) -> dict[str, Any]:
    el = styles_root.find(f"{_w('docDefaults')}/{_w('rPrDefault')}/{_w('rPr')}")
    if el is None:
        return {}
    sz = el.find(_w("sz"))
    out: dict[str, Any] = {"sz_half_points": sz.get(_w("val")) if sz is not None else None}
    out.update(_rfonts_attrs(el.find(_w("rFonts"))))
    return out


def _resolve_effective_run_style(
    doc_root: Any, styles_root: Any, paragraph_index: int, run_index: int
) -> dict[str, Any]:
    """路径 3:direct rPr → 段落样式 basedOn 链 → docDefaults 的有效样式解析。"""
    paras = _list_paragraphs(doc_root)
    p = paras[paragraph_index]
    r_el = p.findall(_w("r"))[run_index]

    direct_rfonts = r_el.find(f"{_w('rPr')}/{_w('rFonts')}")
    direct_sz = r_el.find(f"{_w('rPr')}/{_w('sz')}")
    direct = {
        "rFonts_ascii": direct_rfonts.get(_w("ascii")) if direct_rfonts is not None else None,
        "rFonts_eastAsia": direct_rfonts.get(_w("eastAsia")) if direct_rfonts is not None else None,
        "sz_half_points": direct_sz.get(_w("val")) if direct_sz is not None else None,
    }

    p_style_el = p.find(f"{_w('pPr')}/{_w('pStyle')}")
    p_style_id = p_style_el.get(_w("val")) if p_style_el is not None else "Normal"
    chain = _style_chain(styles_root, p_style_id)
    defaults = _doc_defaults(styles_root)

    # 优先级合并:direct > 链上最近定义 > docDefaults
    def first_non_none(attr: str) -> tuple[Any, str]:
        if direct.get(attr) is not None:
            return direct[attr], "direct"
        for level in chain:
            if level.get("found") and level.get(attr) is not None:
                return level[attr], f"style:{level['styleId']}"
        if defaults.get(attr) is not None:
            return defaults[attr], "docDefaults"
        return None, "none"

    eff_ascii, src_ascii = first_non_none("rFonts_ascii")
    eff_east, src_east = first_non_none("rFonts_eastAsia")
    eff_sz, src_sz = first_non_none("sz_half_points")
    sz_pt = float(eff_sz) / 2 if eff_sz is not None else None

    return {
        "paragraph_style_id": p_style_id,
        "direct_formatting": direct,
        "style_chain": chain,
        "doc_defaults": defaults,
        "effective": {
            "rFonts_ascii": eff_ascii,
            "rFonts_ascii_source": src_ascii,
            "rFonts_eastAsia": eff_east,
            "rFonts_eastAsia_source": src_east,
            "sz_half_points": eff_sz,
            "sz_points": sz_pt,
            "sz_source": src_sz,
        },
    }


def _read_numbering(numbering_root: Any, style_num_pr: dict[str, Any]) -> dict[str, Any]:
    """路径 4:numId → abstractNumId → 指定级别的 numFmt/lvlText。"""
    num_id = style_num_pr.get("numId")
    ilvl = style_num_pr.get("ilvl", "0")
    num_el = numbering_root.find(f"{_w('num')}[@{_w('numId')}='{num_id}']")
    abstract_id = None
    if num_el is not None:
        abstract_el = num_el.find(_w("abstractNumId"))
        abstract_id = abstract_el.get(_w("val")) if abstract_el is not None else None

    lvl_info = None
    if abstract_id is not None:
        abs_el = numbering_root.find(f"{_w('abstractNum')}[@{_w('abstractNumId')}='{abstract_id}']")
        if abs_el is not None:
            for lvl in abs_el.findall(_w("lvl")):
                if lvl.get(_w("ilvl")) == ilvl:
                    num_fmt = lvl.find(_w("numFmt"))
                    lvl_text = lvl.find(_w("lvlText"))
                    start = lvl.find(_w("start"))
                    lvl_info = {
                        "ilvl": ilvl,
                        "numFmt": num_fmt.get(_w("val")) if num_fmt is not None else None,
                        "lvlText": lvl_text.get(_w("val")) if lvl_text is not None else None,
                        "start": start.get(_w("val")) if start is not None else None,
                    }
                    break
    return {"numId": num_id, "abstractNumId": abstract_id, "level": lvl_info}


def _read_headers_footers(zf: zipfile.ZipFile, doc_root: Any) -> dict[str, Any]:
    """路径 5:zip 枚举 header/footer part + sectPr 引用 + part 内文本。"""
    names = zf.namelist()
    header_parts = sorted(n for n in names if posixpath.dirname(n) == "word" and "header" in n)
    footer_parts = sorted(n for n in names if posixpath.dirname(n) == "word" and "footer" in n)

    part_texts = {}
    for n in header_parts + footer_parts:
        root = etree.fromstring(zf.read(n))
        part_texts[n] = "".join(t.text or "" for t in root.iter(_w("t")))

    # sectPr 引用(document.xml 中 w:headerReference 的 r:id 与类型)
    refs = []
    for ref in itertools.chain(
        doc_root.iter(_w("headerReference")), doc_root.iter(_w("footerReference"))
    ):
        refs.append(
            {
                "kind": etree.QName(ref).localname,
                "type": ref.get(_w("type")),
                "r_id": ref.get(f"{{{R}}}id"),
            }
        )
    return {
        "header_parts": header_parts,
        "footer_parts": footer_parts,
        "part_texts": part_texts,
        "references": refs,
    }


def run() -> dict[str, Any]:
    """执行探针,返回 {status, findings, details}。"""
    findings: list[str] = []
    docx_bytes = _build_docx_bytes()
    zf = zipfile.ZipFile(io.BytesIO(docx_bytes))
    doc_root = etree.fromstring(zf.read("word/document.xml"))
    styles_root = etree.fromstring(zf.read("word/styles.xml"))

    details: dict[str, Any] = {"zip_parts": zf.namelist()}

    # ---- 路径 1:锚点 ----
    bookmarks = _enum_bookmarks(doc_root)
    details["bookmarks"] = bookmarks
    target = [b for b in bookmarks if b["name"] == BOOKMARK_NAME]
    if not target:
        raise RuntimeError(f"未能在 document.xml 中读到锚点 {BOOKMARK_NAME!r}")
    b = target[0]
    if b["id"] != BOOKMARK_ID or "anchor paragraph text" not in (b["paragraph_text"] or ""):
        raise RuntimeError(f"锚点 id/所在段落不符: {b!r}")
    if not b["has_matching_bookmarkEnd"]:
        raise RuntimeError("bookmarkStart 没有配对的 bookmarkEnd")

    # ---- 路径 2:文本结构 ----
    text_struct = _enum_text_structure(doc_root)
    details["text_structure"] = text_struct
    all_para_text = [
        "".join(r["text"] for r in p["runs"]) for p in text_struct["paragraphs"]
    ]

    def _table_cell_texts(tables: list[dict[str, Any]]):
        for t in tables:
            for row in t["rows"]:
                for cell in row["cells"]:
                    yield from cell

    placeholder_in_paragraphs = any("{{ payload }}" in t for t in all_para_text)
    placeholder_in_tables = any(
        "{{ payload }}" in seg for seg in _table_cell_texts(text_struct["tables"])
    )
    if not (placeholder_in_paragraphs or placeholder_in_tables):
        raise RuntimeError("未能在文本结构中定位表格单元格占位符 {{ payload }}")
    placeholder_cell = None
    for t in text_struct["tables"]:
        for row in t["rows"]:
            for ci, cell in enumerate(row["cells"]):
                if any("{{ payload }}" in seg for seg in cell):
                    placeholder_cell = {"row": row["row"], "col": ci}
    details["placeholder_cell"] = placeholder_cell

    # ---- 路径 3:有效样式(direct run 与 heading run 两个样本)----
    paras = _list_paragraphs(doc_root)
    direct_para_idx = next(
        i for i, p in enumerate(paras) if "direct formatted" in _para_text(p)
    )
    heading_para_idx = next(i for i, p in enumerate(paras) if _para_text(p) == "heading text")
    direct_style = _resolve_effective_run_style(doc_root, styles_root, direct_para_idx, 0)
    heading_style = _resolve_effective_run_style(doc_root, styles_root, heading_para_idx, 0)
    details["effective_style_direct_run"] = direct_style
    details["effective_style_heading_run"] = heading_style

    eff_d = direct_style["effective"]
    if eff_d["rFonts_ascii"] != DIRECT_FONT_ASCII:
        raise RuntimeError(f"direct formatting 字体解析不符: {eff_d!r}")
    if eff_d["rFonts_eastAsia"] != DIRECT_FONT_EAST_ASIA:
        raise RuntimeError(f"direct formatting 中文字体解析不符: {eff_d!r}")
    if eff_d["sz_points"] != DIRECT_SIZE_PT:
        raise RuntimeError(f"direct formatting 字号解析不符: {eff_d!r}")
    eff_h = heading_style["effective"]
    if eff_h["sz_half_points"] in (None, "22"):
        raise RuntimeError(f"Heading 1 有效字号未解析到样式链定义: {eff_h!r}")

    # ---- 路径 4:numbering ----
    # List Bullet 段的样式 numPr:从 styles.xml 的 ListBullet 取 numId
    list_para_idx = next(i for i, p in enumerate(paras) if _para_text(p) == "bullet item")
    list_p = paras[list_para_idx]
    p_style_el = list_p.find(f"{_w('pPr')}/{_w('pStyle')}")
    list_style_id = p_style_el.get(_w("val")) if p_style_el is not None else None
    style_el = styles_root.find(f"{_w('style')}[@{_w('styleId')}='{list_style_id}']")
    style_numpr = style_el.find(f"{_w('pPr')}/{_w('numPr')}")
    num_id_el = style_numpr.find(_w("numId"))
    ilvl_el = style_numpr.find(_w("ilvl"))
    style_num_pr = {
        "numId": num_id_el.get(_w("val")) if num_id_el is not None else None,
        "ilvl": ilvl_el.get(_w("val")) if ilvl_el is not None else "0",
    }
    numbering = _read_numbering(
        etree.fromstring(zf.read("word/numbering.xml")), style_num_pr
    )
    details["list_style_id"] = list_style_id
    details["style_num_pr"] = style_num_pr
    details["numbering"] = numbering
    if numbering["abstractNumId"] is None or numbering["level"] is None:
        raise RuntimeError(f"numbering 解析失败: {numbering!r}")

    # ---- 路径 5:页眉页脚 ----
    hf = _read_headers_footers(zf, doc_root)
    details["headers_footers"] = hf
    if "word/header1.xml" not in hf["header_parts"] or "word/footer1.xml" not in hf["footer_parts"]:
        raise RuntimeError(f"页眉页脚 part 枚举不符: {hf!r}")
    if hf["part_texts"].get("word/header1.xml") != "HEADER TEXT":
        raise RuntimeError(f"页眉文本读取不符: {hf['part_texts']!r}")
    if not hf["references"]:
        raise RuntimeError("sectPr 中未读到 headerReference/footerReference")

    # ---- python-docx 高层 API 能力对照 ----
    probe_doc = Document(io.BytesIO(docx_bytes))
    high_level = {
        "has_bookmark_api": any("bookmark" in n.lower() for n in dir(probe_doc)) or any(
            "bookmark" in n.lower() for p in probe_doc.paragraphs for n in dir(p)
        ),
        "has_numbering_attr": hasattr(probe_doc, "numbering_part")
        or hasattr(probe_doc, "numbering"),
        "header_via_api": probe_doc.sections[0].header.paragraphs[0].text,
        "table_cell_via_api": probe_doc.tables[0].cell(1, 1).paragraphs[0].text,
        "style_font_via_api": probe_doc.paragraphs[heading_para_idx]
        .runs[0]
        .font.size,  # 高层 API 只看 direct formatting,样式级需 style.font
    }
    details["python_docx_high_level_access"] = high_level

    # ---- findings ----
    findings.append(
        f"锚点:python-docx 没有任何公开 bookmark API(dir 检查={high_level['has_bookmark_api']}),"
        "写入需底层 lxml(run._r.addprevious(bookmarkStart)),读取用 lxml 枚举 "
        "w:bookmarkStart[@w:name/@w:id] + 沿祖先定位 w:p 完全可行;配对 bookmarkEnd 按 w:id 匹配。"
    )
    findings.append(
        "文本结构:lxml 按 w:body→w:p→w:r→{w:t,w:tab,w:br} 枚举可行;表格内段落同样被 "
        "body.iter(w:p) 覆盖(占位符段落实测位于 table row=1 col=1 单元格),"
        "cell 坐标需走 w:tbl→w:tr→w:tc 显式遍历。python-docx 高层 API 读表格文本可行"
        f"(实测 cell(1,1)={high_level['table_cell_via_api']!r})。"
    )
    findings.append(
        f"有效样式解析(direct 样本):direct w:rPr 优先——ascii={eff_d['rFonts_ascii']!r}、"
        f"eastAsia={eff_d['rFonts_eastAsia']!r}、字号 {eff_d['sz_points']}pt(w:sz=28 半磅)"
        "全部来自 direct formatting,解析路径可行。注意 python-docx 高层 run.font.size 只反映 "
        f"direct formatting(实测 heading run 的 run.font.size="
        f"{high_level['style_font_via_api']!r}),"
        "不解析样式链——有效样式必须自实现 direct→basedOn 链→docDefaults 合并。"
    )
    findings.append(
        f"有效样式解析(Heading 1 样本):实测字号 {eff_h['sz_half_points']} 半磅"
        f"({eff_h['sz_points']}pt,来源 {eff_h['sz_source']}),basedOn 链 "
        f"{[lvl['styleId'] for lvl in heading_style['style_chain']]} 可完整走通到 docDefaults。"
    )
    dd = heading_style["doc_defaults"]
    findings.append(
        "python-docx 默认模板的字体大量走【主题间接引用】:docDefaults 的 w:rFonts 只有 "
        f"asciiTheme={dd.get('rFonts_asciiTheme')!r}/eastAsiaTheme={dd.get('rFonts_eastAsiaTheme')!r},"
        "没有字面 w:ascii/w:eastAsia 属性——只读 w:rFonts@w:ascii 会得到 None;"
        "P1 的字体解析必须同时处理 *Theme 属性(必要时到 theme1.xml 的 minorFont/majorFont 解析)。"
        f"docDefaults 字号实测 {dd.get('sz_half_points')} 半磅(=11pt)。"
    )
    findings.append(
        f"numbering:可行但全靠 lxml——python-docx 无公开 numbering API(hasattr 检查="
        f"{high_level['has_numbering_attr']})。链路:段落/样式 w:pPr/w:numPr(ListBullet 样式 "
        f"numId={style_num_pr['numId']}、无 w:ilvl 元素时按 0 级处理)→ numbering.xml "
        f"w:num[@numId]→w:abstractNumId(={numbering['abstractNumId']})→ "
        f"w:abstractNum/w:lvl[@w:ilvl] 取 numFmt={numbering['level']['numFmt']!r}、"
        f"lvlText={numbering['level']['lvlText']!r}。默认模板自带 word/numbering.xml。"
    )
    findings.append(
        f"页眉页脚:zip 枚举 word/header*.xml/footer*.xml 可行(实测 parts={hf['header_parts']} + "
        f"{hf['footer_parts']});part 内文本用同样的 w:t 枚举读取(header1.xml 实测="
        f"{hf['part_texts']['word/header1.xml']!r});sectPr 的 w:headerReference 通过 r:id 关联,"
        "要映射到具体 part 需读 word/_rels/document.xml.rels。python-docx 高层 API 读页眉文本可行"
        f"(实测={high_level['header_via_api']!r})。"
    )
    findings.append(
        "结论:五条读取路径全部可行;其中【锚点、numbering、有效样式合并(含 basedOn/docDefaults/"
        "theme 字体)】必须直接 lxml 解析 zip 内 XML,python-docx 高层 API 仅覆盖普通段落/表格文本/"
        "页眉页脚文本的读取。"
    )

    return {
        "status": "pass",
        "probe": "probe4_ooxml_anchors_styles",
        "findings": findings,
        "details": details,
    }


if __name__ == "__main__":
    import json
    import sys

    try:
        report = run()
    except Exception as exc:  # 探针本身跑不通 → fail + 退出码 1
        import traceback

        report = {
            "status": "fail",
            "probe": "probe4_ooxml_anchors_styles",
            "error": f"{type(exc).__module__}.{type(exc).__name__}: {exc}",
            "traceback": traceback.format_exc(),
        }
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    sys.exit(0 if report["status"] == "pass" else 1)
