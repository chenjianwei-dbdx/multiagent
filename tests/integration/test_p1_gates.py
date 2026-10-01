"""P1 gate bite tests: injected static text (T13), override rendering,
multi-span paragraphs — over the real renderer + extractor stack."""

from __future__ import annotations

from datetime import UTC, datetime

from omas.core.digest import payload_digest
from omas.domain.decisions import AwaitingEvent, Decision, SlotOverride
from omas.domain.findings import GateStatus
from omas.domain.ids import (
    AwaitingEventId,
    DecisionId,
    TaskId,
    new_awaiting_event_id,
    new_decision_id,
)
from omas.gates.gate_b import ProvenanceGate
from omas.pipeline.deterministic import DeterministicRenderPipeline
from omas.pipeline.render_ir_builder import RenderIRBuilder
from omas.renderers.docx_renderer import render_docx
from tests.integration.test_p1_closed_loop import build_case, run_pipeline


def _render_clean(env, weekly, case):
    task, version_id, contract, _material, binding, binding_artifact, _texts = case
    pipeline = DeterministicRenderPipeline(env.home, env.store, env.ledger)
    resolver = pipeline._make_span_resolver()
    ir = RenderIRBuilder(env.ledger).build(
        task_id=TaskId(task.task_id),
        epoch=task.epoch,
        contract=contract,
        template_version_id=version_id,
        template_docx_artifact_id=binding_artifact.artifact_id,
        binding=binding,
        plan_artifact_id=binding_artifact.artifact_id,
        binding_artifact_id=binding_artifact.artifact_id,
    )
    return render_docx(
        ir=ir, contract=contract, template_docx=weekly.docx_bytes, resolve_span=resolver
    ), ir, resolver


def _inject_paragraph(docx_bytes: bytes, text: str) -> bytes:
    import io

    import docx as python_docx

    document = python_docx.Document(io.BytesIO(docx_bytes))
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_injected_static_text_is_blocked(loop_env, weekly) -> None:
    """T13: text added outside any slot fails Gate B."""
    case = build_case(loop_env, weekly)
    _task, _version_id, contract, _material, _binding, _ba, _texts = case
    clean, ir, resolver = _render_clean(loop_env, weekly, case)

    gate = ProvenanceGate()
    clean_report = gate.run(
        candidate_docx=clean,
        template_docx=weekly.docx_bytes,
        contract=contract,
        ir=ir,
        resolve_span=resolver,
        static_map=weekly.static_map,
    )
    assert clean_report.overall_status() is GateStatus.PASS

    injected = _inject_paragraph(clean, "未经任何来源许可的注入文字")
    injected_report = gate.run(
        candidate_docx=injected,
        template_docx=weekly.docx_bytes,
        contract=contract,
        ir=ir,
        resolve_span=resolver,
        static_map=weekly.static_map,
    )
    assert injected_report.overall_status() is GateStatus.FAIL
    sequence = next(r for r in injected_report.results if r.check_id == "content_sequence")
    assert "extra paragraph" in (sequence.actual or "")


def test_override_slot_renders_empty_and_delivers(loop_env, weekly) -> None:
    """risks is allow_user_omit in the fixture: a valid decision yields an
    empty paragraph, gates pass, and the manifest discloses the override."""
    from omas.domain.ir import MissingBinding

    extra = [MissingBinding(binding_status="missing", slot_id="risks", reason_code="no_material")]
    case = build_case(loop_env, weekly, drop_slots={"risks"}, extra_bindings=extra)
    task, _version_id, _contract, material, _binding, _binding_artifact, _texts = case

    awaiting = AwaitingEvent(
        awaiting_event_id=new_awaiting_event_id(),
        task_id=TaskId(task.task_id),
        epoch=task.epoch,
        missing_slot_ids=("risks",),
        created_at=datetime.now(UTC),
    )
    loop_env.ledger.awaiting_events.create(awaiting)
    decision = Decision(
        decision_id=new_decision_id(),
        task_id=TaskId(task.task_id),
        expected_epoch=task.epoch,
        awaiting_event_id=AwaitingEventId(awaiting.awaiting_event_id),
        action="omit_slot",
        payload_digest=payload_digest("omit", "risks"),
        created_at=datetime.now(UTC),
    )
    loop_env.ledger.decisions.accept(decision)
    loop_env.ledger.slot_overrides.insert(
        f"ovr_{decision.decision_id.removeprefix('dec_')}",
        SlotOverride(
            task_id=TaskId(task.task_id),
            slot_id="risks",
            decision_id=DecisionId(decision.decision_id),
            epoch=task.epoch,
            reason_artifact_id=material.artifact_id,
        ),
    )

    delivery = run_pipeline(loop_env, weekly, case)
    manifest_artifact = loop_env.ledger.artifacts.get(delivery.manifest_artifact_id)
    assert manifest_artifact is not None
    manifest = loop_env.store.read_verified(
        manifest_artifact.relative_path, manifest_artifact.sha256
    )
    assert b'"overrides"' in manifest


def test_multi_span_slot_renders_separate_paragraphs(loop_env, weekly) -> None:
    """sales_summary binds two section paragraphs → two adjacent paragraphs."""
    case = build_case(loop_env, weekly)
    sales_texts = case[6]["sales_summary"]
    assert len(sales_texts) == 2  # the fixture's sales section has two paragraphs
    clean, ir, resolver = _render_clean(loop_env, weekly, case)

    from omas.docx.extract import extract_details, paragraph_plain_text

    document = extract_details(clean).part("document")
    texts = [paragraph_plain_text(d.paragraph) for d in document.paragraphs]
    for text in sales_texts:
        assert text in texts

    gate = ProvenanceGate()
    report = gate.run(
        candidate_docx=clean,
        template_docx=weekly.docx_bytes,
        contract=case[2],
        ir=ir,
        resolve_span=resolver,
        static_map=weekly.static_map,
    )
    assert report.overall_status() is GateStatus.PASS
