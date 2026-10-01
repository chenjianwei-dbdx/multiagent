"""DOCX renderer unit tests.

Covers (task contract): code-point fidelity (CJK, XML specials, spaces,
Tab/LF, emoji incl. ZWJ), multi-span -> multi-paragraph, single span,
override -> empty paragraph, template hash gate, IR-vs-contract slot gate,
resolver error pass-through, leftover-placeholder defense, static-region
immutability, autoescape, and the class wrapper.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from omas.domain.errors import ArtifactHashMismatchError, SpanHashMismatchError
from omas.domain.ids import new_artifact_id
from omas.domain.ir import RenderSlot
from omas.domain.spans import SourceSpanRef
from omas.renderers import DocxRenderer, RenderError, render_docx

# Every payload must round-trip exactly (code point for code point), with
# Tab/LF surviving through w:tab/w:br and no trimming or collapsing.
FIDELITY_PAYLOADS = (
    "中文测试，包含全角标点。",
    'a<b>c & d"e\'f>g',
    "<script>alert('x' & \"y\")</script>",
    "a  b   c",
    "  padded  ",
    "a\tb",
    "line1\nline2",
    "😀 和 👨‍👩‍👧‍👦",
    "mixed: 中<文>\t&\n\"😀\"  空格",
)


@pytest.mark.parametrize("payload", FIDELITY_PAYLOADS)
def test_single_span_roundtrip_is_codepoint_exact(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
    document_xml,
    payload,
) -> None:
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans(payload)
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, [payload])

    out = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    rendered = paragraph_texts(out)
    template_paragraphs = paragraph_texts(template)
    assert rendered[:2] == template_paragraphs[:2]  # static head untouched
    assert rendered[2] == payload  # exact, no trim / collapse / escaping loss
    assert "{{" not in document_xml(out)
    assert resolver.calls == list(spans)  # resolved exactly once


def test_output_is_a_valid_docx_for_python_docx(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
) -> None:
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans("普通文本。")
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, ["普通文本。"])

    out = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    doc = Document(io.BytesIO(out))  # must open as a normal DOCX package
    body_texts = [p.text for p in doc.paragraphs]
    assert body_texts[0] == "项目周报"
    assert body_texts[-1] == "普通文本。"


def test_autoescape_keeps_markup_as_literal_text(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
) -> None:
    payload = '<script>alert("x & y")</script>'
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans(payload)
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, [payload])

    out = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    assert paragraph_texts(out)[2] == "<script>alert(\"x & y\")</script>"


def test_multi_span_renders_one_paragraph_per_span_in_order(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
    document_xml,
) -> None:
    texts = ("第一段：销售达标。", "第二段：库存有所下降。", "第三段：回款节奏正常。")
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans(*texts)
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, list(texts))

    out = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    rendered = paragraph_texts(out)
    assert rendered[2:5] == list(texts)  # each span is its own paragraph, in order
    assert len(rendered) == 5  # 2 static + 3 span paragraphs; nothing else added
    assert resolver.calls == list(spans)  # every span resolved exactly once, IR order
    assert "{{" not in document_xml(out)


def test_override_slot_renders_empty_paragraph_without_filler(
    make_template_docx,
    make_contract,
    make_ir,
    resolver,
    paragraph_texts,
    document_xml,
) -> None:
    template = make_template_docx(("next_plan",))
    contract = make_contract(template, ("next_plan",), allow_user_omit=("next_plan",))
    ir = make_ir(
        [
            RenderSlot(
                slot_id="next_plan",
                spans=(),
                style_key="body",
                override_artifact_id=new_artifact_id(),
            )
        ]
    )

    out = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    assert paragraph_texts(out)[2] == ""  # empty paragraph kept, no text inserted
    xml = document_xml(out)
    assert "{{" not in xml
    assert "暂无" not in xml
    assert resolver.calls == []  # override resolves nothing


def test_template_hash_mismatch_raises(
    template_docx,
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
) -> None:
    # Contract whose docx_sha256 belongs to a different template package.
    other_contract = make_contract(make_template_docx(("risks",)))
    spans = make_spans("文本")
    ir = make_ir([RenderSlot(slot_id="sales_summary", spans=spans, style_key="body")])

    with pytest.raises(ArtifactHashMismatchError) as excinfo:
        render_docx(
            ir=ir,
            contract=other_contract,
            template_docx=template_docx,
            resolve_span=lambda _ref: "never reached",
        )
    assert excinfo.value.code == "ARTIFACT_HASH_MISMATCH"


def test_ir_slot_outside_contract_raises_value_error(
    template_docx,
    contract,
    make_spans,
    make_ir,
) -> None:
    spans = make_spans("文本")
    ir = make_ir([RenderSlot(slot_id="ghost", spans=spans, style_key="body")])

    with pytest.raises(ValueError, match="ghost"):
        render_docx(
            ir=ir,
            contract=contract,
            template_docx=template_docx,
            resolve_span=lambda _ref: "never reached",
        )


def test_render_error_is_the_domain_error() -> None:
    # Single source of truth: gates/pipeline catch omas.domain.errors.RenderError.
    from omas.domain.errors import RenderError as DomainRenderError

    assert RenderError is DomainRenderError
    assert RenderError.code == "RENDER_FAILED"


def test_resolver_failure_propagates_unwrapped(
    template_docx,
    contract,
    make_spans,
    make_ir,
) -> None:
    spans = make_spans("文本")
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])

    def failing(_ref: SourceSpanRef) -> str:
        raise SpanHashMismatchError("simulated span verification failure")

    with pytest.raises(SpanHashMismatchError) as excinfo:
        render_docx(
            ir=ir, contract=contract, template_docx=template_docx, resolve_span=failing
        )
    assert type(excinfo.value) is SpanHashMismatchError  # not wrapped, not swallowed
    assert excinfo.value.code == "SPAN_HASH_MISMATCH"
    assert not isinstance(excinfo.value, RenderError)


def test_uncovered_template_placeholder_raises_render_error(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
) -> None:
    template = make_template_docx(extra_placeholders=("{{ mystery }}",))
    contract = make_contract(template)
    spans = make_spans("a", "b", "c")
    ir = make_ir(
        [
            RenderSlot(slot_id="sales_summary", spans=spans[0:1], style_key="body"),
            RenderSlot(slot_id="risks", spans=spans[1:2], style_key="body"),
            RenderSlot(slot_id="next_plan", spans=spans[2:3], style_key="body"),
        ]
    )
    resolver.given(spans, ["a", "b", "c"])

    with pytest.raises(RenderError) as excinfo:
        render_docx(
            ir=ir, contract=contract, template_docx=template, resolve_span=resolver
        )
    assert excinfo.value.code == "RENDER_FAILED"


def test_span_text_containing_placeholder_syntax_raises_render_error(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
) -> None:
    payload = "备注 {{ note }} 需人工复核"
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans(payload)
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, [payload])

    with pytest.raises(RenderError) as excinfo:
        render_docx(
            ir=ir, contract=contract, template_docx=template, resolve_span=resolver
        )
    assert excinfo.value.code == "RENDER_FAILED"


def test_static_paragraphs_and_table_unchanged(
    template_docx,
    contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
    table_texts,
    document_xml,
) -> None:
    assert paragraph_texts(template_docx)[0] == "项目周报"  # CJK static title present
    template_paragraphs = paragraph_texts(template_docx)
    template_tables = table_texts(template_docx)

    spans = make_spans("销售摘要文本。", "风险描述文本。", "下周计划文本。")
    ir = make_ir(
        [
            RenderSlot(slot_id="sales_summary", spans=spans[0:1], style_key="body"),
            RenderSlot(slot_id="risks", spans=spans[1:2], style_key="body"),
            RenderSlot(slot_id="next_plan", spans=spans[2:3], style_key="body"),
        ]
    )
    resolver.given(spans, ["销售摘要文本。", "风险描述文本。", "下周计划文本。"])

    out = render_docx(
        ir=ir, contract=contract, template_docx=template_docx, resolve_span=resolver
    )

    out_paragraphs = paragraph_texts(out)
    assert out_paragraphs[:2] == template_paragraphs[:2]
    assert table_texts(out) == template_tables
    assert "{{" not in document_xml(out)


def test_mixed_single_multi_and_override_layout(
    template_docx,
    contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
    document_xml,
) -> None:
    sales = "销售整体达标，华东区超额完成。"
    risk_texts = ("风险一：库存周转偏低。", "风险二：两笔回款延期。", "风险三：人员流动。")
    spans = make_spans(sales, *risk_texts)
    ir = make_ir(
        [
            RenderSlot(slot_id="sales_summary", spans=spans[0:1], style_key="body"),
            RenderSlot(slot_id="risks", spans=spans[1:4], style_key="body"),
            RenderSlot(
                slot_id="next_plan",
                spans=(),
                style_key="body",
                override_artifact_id=new_artifact_id(),
            ),
        ]
    )
    resolver.given(spans, [sales, *risk_texts])

    out = render_docx(
        ir=ir, contract=contract, template_docx=template_docx, resolve_span=resolver
    )

    expected = [*paragraph_texts(template_docx)[:2], sales, *risk_texts, ""]
    assert paragraph_texts(out) == expected
    assert resolver.calls == list(spans)  # IR order: sales span, then risk spans
    assert "{{" not in document_xml(out)


def test_renderer_class_matches_module_function(
    make_template_docx,
    make_contract,
    make_spans,
    make_ir,
    resolver,
    paragraph_texts,
    document_xml,
) -> None:
    template = make_template_docx(("risks",))
    contract = make_contract(template, ("risks",), allow_user_omit=())
    spans = make_spans("类封装输出。")
    ir = make_ir([RenderSlot(slot_id="risks", spans=spans, style_key="body")])
    resolver.given(spans, ["类封装输出。"])

    via_class = DocxRenderer().render(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )
    via_function = render_docx(
        ir=ir, contract=contract, template_docx=template, resolve_span=resolver
    )

    assert paragraph_texts(via_class) == paragraph_texts(via_function)
    assert "{{" not in document_xml(via_class)
