"""Chat console JSON API（ADR 0002 amendment）：会话/消息/进度流/应答/下载/模板库。

全部 local_only 确定性路径，后台 TaskRunner 为真实线程——与 SPA 相同的轮询方式等待终态。
"""

from __future__ import annotations

import io
import time

import docx
from tests.unit.web.conftest import MATERIAL_PARTIAL, MATERIAL_RISKS


def _poll_terminal(web_env, task_id: str, timeout: float = 30.0) -> dict:
    """与 SPA 相同的 feed 轮询，直到终态；返回最后的 view。"""
    after = 0
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = web_env.client.get(f"/api/tasks/{task_id}/feed?after={after}")
        assert response.status_code == 200, response.text
        payload = response.json()
        after = payload["last_seq"]
        status = payload["view"]["status"]
        if status in ("completed", "failed", "cancelled"):
            return payload["view"]
        if status == "awaiting_user":
            return payload["view"]
        time.sleep(0.1)
    raise AssertionError(f"task {task_id} did not reach a terminal state in {timeout}s")


def _make_conversation(web_env, data_policy: str = "local_only") -> str:
    response = web_env.client.post(
        "/api/conversations",
        json={
            "template_version_id": web_env.version_id,
            "title": "测试会话",
            "data_policy": data_policy,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["conversation_id"]


def _send_message(web_env, conversation_id: str, intent: str, materials: list[tuple[str, str]]):
    """Send a chat turn; wait for triage (local_only: direct) and return the task id."""
    import time as _time

    files = [
        ("materials", (name, text.encode("utf-8"), "text/markdown"))
        for name, text in materials
    ]
    response = web_env.client.post(
        f"/api/conversations/{conversation_id}/messages",
        data={"intent": intent},
        files=files,
    )
    assert response.status_code == 202, response.text
    deadline = _time.monotonic() + 30
    while _time.monotonic() < deadline:
        detail = web_env.client.get(f"/api/conversations/{conversation_id}").json()
        tasks = [m for m in detail["messages"] if m["kind"] == "task_started"]
        if tasks and not detail["busy"]:
            return tasks[-1]["task_id"]
        _time.sleep(0.1)
    raise AssertionError("turn did not produce a task in time")


def test_conversation_requires_registered_template(web_env) -> None:
    response = web_env.client.post(
        "/api/conversations",
        json={"template_version_id": "tver_" + "9" * 32, "data_policy": "local_only"},
    )
    assert response.status_code == 404
    assert "not registered" in response.text


def test_message_completes_and_downloads(web_env) -> None:
    conversation_id = _make_conversation(web_env)
    task_id = _send_message(
        web_env, conversation_id, "生成本周周报", [("week.md", web_env.full_material)]
    )
    view = _poll_terminal(web_env, task_id)
    assert view["status"] == "completed"
    assert view["delivery"] is not None

    download = web_env.client.get(f"/api/tasks/{task_id}/download")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    detail = web_env.client.get(f"/api/conversations/{conversation_id}").json()
    kinds = [m["kind"] for m in detail["messages"]]
    assert "text" in kinds and "task_started" in kinds
    assert any(t["task_id"] == task_id and t["status"] == "completed" for t in detail["tasks"])


def test_gap_then_respond_via_api(web_env) -> None:
    conversation_id = _make_conversation(web_env)
    task_id = _send_message(
        web_env, conversation_id, "生成周报", [("week.md", MATERIAL_PARTIAL)]
    )
    view = _poll_terminal(web_env, task_id)
    assert view["status"] == "awaiting_user"
    assert "risks" in view["awaiting"]["missing_slot_ids"]

    respond = web_env.client.post(
        f"/api/tasks/{task_id}/respond",
        files=[("materials", ("risks.md", MATERIAL_RISKS.encode("utf-8"), "text/markdown"))],
    )
    assert respond.status_code == 202, respond.text
    view = _poll_terminal(web_env, task_id)
    assert view["status"] == "completed"
    count = web_env.container.ledger.connection.execute(
        "SELECT COUNT(*) AS n FROM deliveries WHERE task_id = ?", (task_id,)
    ).fetchone()["n"]
    assert count == 1


def test_feed_includes_tool_and_step_events(web_env) -> None:
    conversation_id = _make_conversation(web_env)
    task_id = _send_message(
        web_env, conversation_id, "生成周报", [("week.md", web_env.full_material)]
    )
    _poll_terminal(web_env, task_id)
    feed = web_env.client.get(f"/api/tasks/{task_id}/feed?after=0").json()
    codes = [e["event_code"] for e in feed["events"]]
    assert "task_submitted" in codes
    assert "inventory" in codes


def _plain_template_docx(slot_names: list[str]) -> bytes:
    """带占位符但无书签的普通文档（模拟用户上传）。"""
    document = docx.Document()
    document.add_paragraph("值班交接单", style="Title")
    for slot in slot_names:
        document.add_paragraph(slot.replace("_", "事件"), style="Heading 2")
        document.add_paragraph("{{ " + slot + " }}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_template_upload_scaffold_and_rename(web_env) -> None:
    docx_bytes = _plain_template_docx(["handover_items", "notes"])
    precheck = web_env.client.post(
        "/api/templates/precheck",
        files=[("docx", ("handover.docx", docx_bytes, "application/octet-stream"))],
    )
    assert precheck.status_code == 200
    assert precheck.json()["slots"] == ["handover_items", "notes"]

    import json as jsonlib

    upload = web_env.client.post(
        "/api/templates",
        data={
            "template_id": "handover-note",
            "display_name": "值班交接单",
            "description": "交接班时填写事项与备注",
            "slot_semantics": jsonlib.dumps(
                {"handover_items": "交接事项清单", "notes": "备注"}, ensure_ascii=False
            ),
        },
        files=[("docx", ("handover.docx", docx_bytes, "application/octet-stream"))],
    )
    assert upload.status_code == 201, upload.text
    body = upload.json()
    assert body["activatable"] is True, body["findings"]
    assert body["slots"] == ["handover_items", "notes"]

    templates = web_env.client.get("/api/templates").json()["templates"]
    mine = next(t for t in templates if t["template_id"] == "handover-note")
    assert mine["display_name"] == "值班交接单"
    assert mine["description"] == "交接班时填写事项与备注"
    assert {s["slot_id"] for s in mine["slots"]} == {"handover_items", "notes"}

    rename = web_env.client.post(
        "/api/templates/handover-note/meta",
        json={"display_name": "交接单 v2", "description": "更新后的简介"},
    )
    assert rename.status_code == 200
    templates = web_env.client.get("/api/templates").json()["templates"]
    mine = next(t for t in templates if t["template_id"] == "handover-note")
    assert mine["display_name"] == "交接单 v2"


def test_template_upload_without_placeholders_rejected(web_env) -> None:
    document = docx.Document()
    document.add_paragraph("没有占位符的普通文档")
    buffer = io.BytesIO()
    document.save(buffer)
    response = web_env.client.post(
        "/api/templates/precheck",
        files=[("docx", ("plain.docx", buffer.getvalue(), "application/octet-stream"))],
    )
    assert response.status_code == 200
    assert response.json()["slots"] == []


def _qa_turn_app(web_env, model_fn):
    """Replace the turn executor with one using a scripted offline model."""
    from omas.config.settings import ModelsConfig, ModelSettings
    from omas.web.turns import TurnExecutor

    web_env.client.app.state.turns = TurnExecutor(
        web_env.home,
        ModelsConfig(
            model=ModelSettings(
                provider="anthropic", model_id="test", base_url="http://127.0.0.1:9"
            )
        ),
        model_builder=lambda settings, gateway: model_fn(),
    )
    return web_env


def _make_llm_conversation(web_env) -> str:
    response = web_env.client.post(
        "/api/conversations",
        json={
            "template_version_id": web_env.version_id,
            "title": "问答会话",
            "data_policy": "llm_allowed",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["conversation_id"]


def _wait_turn(web_env, conversation_id: str, timeout: float = 20.0) -> dict:
    import time

    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        detail = web_env.client.get(f"/api/conversations/{conversation_id}").json()
        if not detail["busy"] and any(
            m["kind"] in ("answer", "clarify", "task_started", "note") for m in detail["messages"]
        ):
            return detail
        time.sleep(0.1)
    raise AssertionError("turn did not finish")


def test_question_turn_answers_via_qa(web_env) -> None:
    from pydantic_ai.messages import ModelMessage, ModelResponse, ToolCallPart
    from pydantic_ai.models.function import AgentInfo, FunctionModel


    def triage_fn(messages: list[ModelMessage], info: AgentInfo) -> ModelResponse:
        return ModelResponse(
            parts=[
                ToolCallPart(
                    tool_name=info.output_tools[0].name,
                    args={"category": "question", "reason": "提问"},
                )
            ]
        )

    triage_model = FunctionModel(triage_fn)
    _qa_turn_app(web_env, lambda: triage_model)
    def qa_fn(messages: list[ModelMessage], info: AgentInfo) -> ModelResponse:
        return ModelResponse(
            parts=[
                ToolCallPart(
                    tool_name=info.output_tools[0].name,
                    args={"answer": "晴，二十度。", "sources": ["天气 | https://w.e"]},
                )
            ]
        )

    qa_model = FunctionModel(qa_fn)

    from omas.web import turns as turns_mod

    original = turns_mod.TurnExecutor._run_qa

    def patched_run_qa(self, container, row, question):
        from pydantic_ai import Agent

        import omas.agents.qa as qa_mod

        original_agent = qa_mod.QAAgent.answer_question

        def answer_question(self, *, question, search, gateway, agent_factory=None):
            _ = agent_factory  # force the scripted QA model regardless of caller factory

            def factory(**_kw: object) -> object:
                return Agent(model=qa_model, output_type=qa_mod.QAAnswer, retries=1)

            return original_agent(
                self, question=question, search=search, gateway=gateway, agent_factory=factory
            )

        qa_mod.QAAgent.answer_question = answer_question
        try:
            return original(self, container, row, question)
        finally:
            qa_mod.QAAgent.answer_question = original_agent

    turns_mod.TurnExecutor._run_qa = patched_run_qa  # type: ignore[method-assign]
    try:
        conversation_id = _make_llm_conversation(web_env)
        response = web_env.client.post(
            f"/api/conversations/{conversation_id}/messages",
            data={"intent": "今天天气怎么样？"},
        )
        assert response.status_code == 202, response.text
        detail = _wait_turn(web_env, conversation_id)
        answers = [m for m in detail["messages"] if m["kind"] == "answer"]
        assert answers, detail["messages"]
        assert "晴" in answers[-1]["content"]
        assert "https://w.e" in answers[-1]["content"]
        assert not any(m["kind"] == "task_started" for m in detail["messages"])
    finally:
        turns_mod.TurnExecutor._run_qa = original  # type: ignore[method-assign]


def test_needs_info_turn_clarifies(web_env) -> None:
    from pydantic_ai.messages import ModelMessage, ModelResponse, ToolCallPart
    from pydantic_ai.models.function import AgentInfo, FunctionModel

    def triage_fn(messages: list[ModelMessage], info: AgentInfo) -> ModelResponse:
        return ModelResponse(
            parts=[
                ToolCallPart(
                    tool_name=info.output_tools[0].name,
                    args={
                        "category": "needs_info",
                        "reason": "信息不足",
                        "clarifying_question": "你要生成什么文档？请附材料。",
                    },
                )
            ]
        )

    model = FunctionModel(triage_fn)
    _qa_turn_app(web_env, lambda: model)
    conversation_id = _make_llm_conversation(web_env)
    response = web_env.client.post(
        f"/api/conversations/{conversation_id}/messages", data={"intent": "帮我弄个东西"}
    )
    assert response.status_code == 202, response.text
    detail = _wait_turn(web_env, conversation_id)
    clarifies = [m for m in detail["messages"] if m["kind"] == "clarify"]
    assert clarifies and "材料" in clarifies[-1]["content"]
